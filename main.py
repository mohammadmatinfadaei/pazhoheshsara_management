from database import (
    db,
    app,
    Books,
    Admins,
    Actions,
    Students,
    Borrowing,
    render_template,
    redirect,
    request,
    flash,
    make_response,
    jsonify
)

import openpyxl
import os

import jdatetime
from datetime import date, datetime, timedelta


# =========================================================
# مسیر فایل‌های Excel
# =========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

BOOKS_EXCEL = os.path.join(
    BASE_DIR,
    "Books.xlsx"
)

STUDENTS_EXCEL = os.path.join(
    BASE_DIR,
    "Students.xlsx"
)


# =========================================================
# ابزارهای کمکی
# =========================================================




@app.template_filter("jalali")
def jalali_filter(value):

    if value is None:
        return "-"

    try:

        # اگر datetime واقعی باشد
        if isinstance(value, datetime):

            jalali_date = jdatetime.date.fromgregorian(
                date=value.date()
            )

            return jalali_date.strftime("%Y/%m/%d")


        # اگر date واقعی باشد
        if isinstance(value, date):

            jalali_date = jdatetime.date.fromgregorian(
                date=value
            )

            return jalali_date.strftime("%Y/%m/%d")


        # تبدیل به رشته
        value = str(value).strip()

        if not value or value.lower() == "none":
            return "-"


        # پیدا کردن تاریخ میلادی داخل رشته
        import re

        match = re.search(
            r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})",
            value
        )

        if not match:
            return value


        year = int(match.group(1))
        month = int(match.group(2))
        day = int(match.group(3))


        # تبدیل میلادی به جلالی
        gregorian_date = date(
            year,
            month,
            day
        )

        jalali_date = jdatetime.date.fromgregorian(
            date=gregorian_date
        )


        # فقط تاریخ، بدون ساعت
        return (
            f"{jalali_date.year:04d}/"
            f"{jalali_date.month:02d}/"
            f"{jalali_date.day:02d}"
        )


    except Exception:

        return str(value)  
    
    
def get_current_admin():

    user = request.cookies.get("user")

    if not user:
        return None

    return Admins.query.filter_by(
        user_name=user
    ).first()


def parse_datetime(value):

    if not value:
        return None

    try:
        return datetime.strptime(
            str(value),
            "%Y-%m-%d %H:%M:%S"
        )

    except (ValueError, TypeError):

        return None


def parse_date(value):

    if not value:
        return None

    try:
        return datetime.strptime(
            str(value),
            "%Y-%m-%d"
        ).date()

    except (ValueError, TypeError):

        return None


# =========================================================
# بررسی انقضای دانش‌آموزان
# =========================================================

def update_expired_students():

    today = date.today()

    students = Students.query.all()

    changed = False

    for student in students:

        date_expired = False

        # -------------------------------------------------
        # بررسی تاریخ انقضای کارت
        # -------------------------------------------------

        expiration_date = parse_date(
            student.expiration_date
        )

        if expiration_date:

            if expiration_date <= today:

                date_expired = True

        # -------------------------------------------------
        # بررسی فرصت‌ها
        # -------------------------------------------------

        opportunity_expired = (

            student.opportunity is not None

            and

            student.opportunity <= 0

        )

        # -------------------------------------------------
        # تعیین وضعیت نهایی
        # -------------------------------------------------

        new_expired_value = (

            date_expired

            or

            opportunity_expired

        )

        if student.expired != new_expired_value:

            student.expired = new_expired_value

            changed = True

    if changed:

        db.session.commit()


# =========================================================
# بررسی دیرکرد همه امانت‌ها
#
# هر امانت فقط یک بار opportunity را کم می‌کند.
# =========================================================

def update_overdue_students():

    now = datetime.now()

    borrowings = Borrowing.query.all()

    changed = False

    for borrowing in borrowings:

        # -------------------------------------------------
        # اگر قبلاً جریمه شده، دوباره جریمه نشود
        # -------------------------------------------------

        if borrowing.late_penalty_applied:

            continue

        # -------------------------------------------------
        # تاریخ بازگشت
        # -------------------------------------------------

        return_date = parse_datetime(
            borrowing.return_date
        )

        if not return_date:

            continue

        # -------------------------------------------------
        # اگر هنوز موعد بازگشت نرسیده
        # -------------------------------------------------

        if return_date >= now:

            continue

        # -------------------------------------------------
        # پیدا کردن دانش‌آموز
        # -------------------------------------------------

        student = Students.query.filter_by(

            student_id_code=
                borrowing.student_id_code_borrow

        ).first()

        if not student:

            continue

        # -------------------------------------------------
        # اگر opportunity خالی بود
        # -------------------------------------------------

        if student.opportunity is None:

            student.opportunity = 0

        # -------------------------------------------------
        # کسر دقیقاً یک فرصت
        # -------------------------------------------------

        if student.opportunity > 0:

            student.opportunity -= 1

        # -------------------------------------------------
        # علامت‌گذاری امانت به عنوان جریمه‌شده
        # -------------------------------------------------

        borrowing.late_penalty_applied = True

        changed = True

        # -------------------------------------------------
        # اگر فرصت‌ها تمام شدند
        # -------------------------------------------------

        if student.opportunity <= 0:

            student.opportunity = 0

            student.expired = True

    if changed:

        db.session.commit()


# =========================================================
# بررسی وضعیت در هر Request
# =========================================================

@app.before_request
def check_students_status():

    try:

        update_expired_students()

        update_overdue_students()

    except Exception:

        # در صورت وجود مشکل در بررسی وضعیت،
        # خود Request اصلی متوقف نشود.
        db.session.rollback()


# =========================================================
# ثبت عملیات ادمین
# =========================================================

def add_action(action_text):

    user = request.cookies.get("user")

    if not user:

        return

    admin = Admins.query.filter_by(
        user_name=user
    ).first()

    if not admin:

        return

    action = Actions(

        actions_taken=action_text,

        date=str(date.today()),

        time=datetime.now().strftime("%H:%M:%S")

    )

    db.session.add(action)


# =========================================================
# CURRENT ADMIN
# =========================================================

@app.context_processor
def inject_current_admin():

    current_admin = get_current_admin()

    return {
        "current_admin": current_admin
    }


# =========================================================
# HOME
# =========================================================

@app.route('/')
def index():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    current_admin = get_current_admin()

    # -----------------------------------------------------
    # تعداد کل کتاب‌ها
    # -----------------------------------------------------

    total_books = Books.query.count()

    # -----------------------------------------------------
    # تعداد کل دانش‌آموزان
    # -----------------------------------------------------

    total_students = Students.query.count()

    # -----------------------------------------------------
    # تعداد کتاب‌های امانت داده شده
    # -----------------------------------------------------

    borrowed_books = Borrowing.query.count()

    # -----------------------------------------------------
    # تعداد کتاب‌های دیرکرد
    # -----------------------------------------------------

    overdue_books = 0

    now = datetime.now()

    borrowings = Borrowing.query.all()

    for borrowing in borrowings:

        return_date = parse_datetime(
            borrowing.return_date
        )

        if return_date and return_date < now:

            overdue_books += 1

    # -----------------------------------------------------
    # 5 فعالیت آخر برای God User
    # -----------------------------------------------------

    latest_actions = []

    if current_admin and current_admin.god_user:

        latest_actions = Actions.query.order_by(
            Actions.id.desc()
        ).limit(5).all()

    return render_template(

        'index.html',

        current_admin=current_admin,

        total_books=total_books,

        total_students=total_students,

        borrowed_books=borrowed_books,

        overdue_books=overdue_books,

        latest_actions=latest_actions

    )


# =========================================================
# LOGIN
# =========================================================

@app.route('/login', methods=['GET', 'POST'])
def login():

    if request.method == 'POST':

        user_name = request.form.get(
            'username'
        )

        password = request.form.get(
            'password'
        )

        admin = Admins.query.filter_by(
            user_name=user_name,
            password=password
        ).first()

        if admin:

            flash(

                "خوش آمدید جناب {} {}".format(
                    admin.first_name,
                    admin.last_name
                ),

                "success"

            )

            response = make_response(
                redirect('/')
            )

            response.set_cookie(
                "user",
                user_name
            )

            return response

        flash(

            "نام کاربری یا رمز ورود اشتباه است",

            "danger"

        )

    return render_template(
        'login.html'
    )


# =========================================================
# REGISTER
# =========================================================

@app.route('/register', methods=['GET', 'POST'])
def Register():

    if request.method == 'POST':

        first_name = request.form.get(
            'first_name'
        )

        last_name = request.form.get(
            'last_name'
        )

        user_name = request.form.get(
            'username'
        )

        password = request.form.get(
            'password'
        )

        existing_admin = Admins.query.filter_by(
            user_name=user_name
        ).first()

        if existing_admin:

            flash(
                "این نام کاربری قبلاً ثبت شده است",
                "danger"
            )

            return redirect('/register')

        admin = Admins(

            user_name=user_name,

            password=password,

            last_name=last_name,

            first_name=first_name

        )

        try:

            db.session.add(admin)

            db.session.commit()

            flash(
                "عضو شدید",
                "success"
            )

            return redirect('/login')

        except Exception as e:

            db.session.rollback()

            flash(
                "خطا در ثبت حساب: " + str(e),
                "danger"
            )

            return redirect('/register')

    return render_template(
        'Register.html'
    )


# =========================================================
# LOGOUT
# =========================================================

@app.route('/logout')
def logout():

    response = make_response(
        redirect('/login')
    )

    response.delete_cookie("user")

    flash(
        "از حسابتان خارج شدید",
        "danger"
    )

    return response


# =========================================================
# SHOW BOOKS
# =========================================================

@app.route("/show_books")
def show_books():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    search = request.args.get(
        "search",
        ""
    ).strip()

    page = request.args.get(
        "page",
        1,
        type=int
    )

    if search:

        books = Books.query.filter(

            (Books.book_title.contains(search)) |

            (Books.author_translator.contains(search)) |

            (Books.congress_number.contains(search)) |

            (Books.shelf_number.contains(search)) |

            (Books.row_number.contains(search))

        ).paginate(

            page=page,

            per_page=30,

            error_out=False

        )

    else:

        books = Books.query.paginate(

            page=page,

            per_page=30,

            error_out=False

        )

    return render_template(

        "books.html",

        books=books,

        search=search

    )


# =========================================================
# ADD BOOK
# =========================================================

@app.route('/add_book', methods=['GET', 'POST'])
def add_book():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    if request.method == 'POST':

        book_title = request.form.get(
            'book_title'
        )

        author_translator = request.form.get(
            'author_translator'
        )

        congress_number = request.form.get(
            'congress_number'
        )

        shelf_number = request.form.get(
            'shelf_number'
        )

        row_number = request.form.get(
            'row_number'
        )

        comprehensive = (

            request.form.get(
                'comprehensive'
            ) == 'True'

        )

        try:

            book = Books(

                book_title=book_title,

                author_translator=author_translator,

                congress_number=congress_number,

                shelf_number=shelf_number,

                row_number=row_number,

                comprehensive=comprehensive,

                borrower_code="None",

                available=False

            )

            db.session.add(book)

            db.session.flush()

            # -------------------------------------------------
            # Excel
            # -------------------------------------------------

            workbook = openpyxl.load_workbook(
                BOOKS_EXCEL
            )

            sheet = workbook.active

            last_real_row = 1

            for r in range(
                2,
                sheet.max_row + 1
            ):

                if any(

                    sheet.cell(
                        row=r,
                        column=c
                    ).value not in (None, "")

                    for c in range(1, 7)

                ):

                    last_real_row = r

            new_row = last_real_row + 1

            values = [

                book_title,

                author_translator,

                congress_number,

                shelf_number,

                row_number,

                comprehensive

            ]

            for column, value in enumerate(

                values,

                start=1

            ):

                sheet.cell(

                    row=new_row,

                    column=column,

                    value=value

                )

            workbook.save(
                BOOKS_EXCEL
            )

            # -------------------------------------------------
            # ثبت عملیات
            # -------------------------------------------------

            admin = Admins.query.filter_by(
                user_name=user
            ).first()

            if admin:

                add_action(

                    "ادمین {} {} کتاب «{}» را اضافه کرد".format(

                        admin.first_name,

                        admin.last_name,

                        book.book_title

                    )

                )

            db.session.commit()

            flash(
                "کتاب با موفقیت اضافه شد",
                "success"
            )

            return redirect(
                '/show_books'
            )

        except Exception as e:

            db.session.rollback()

            flash(
                "خطا در ثبت کتاب: " + str(e),
                "danger"
            )

            return redirect(
                '/add_book'
            )

    return render_template(
        'add_book.html'
    )


# =========================================================
# DELETE BOOK
# =========================================================

@app.route('/delete_book/<int:book_id>')
def delete_book(book_id):

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    book = Books.query.get(book_id)

    if not book:

        flash(
            "کتاب پیدا نشد",
            "danger"
        )

        return redirect('/show_books')

    try:

        # -------------------------------------------------
        # اگر کتاب در امانت است حذف نشود
        # -------------------------------------------------

        if book.available:

            flash(
                "کتاب در حال حاضر امانت داده شده و نمی‌توان آن را حذف کرد.",
                "danger"
            )

            return redirect('/show_books')

        workbook = openpyxl.load_workbook(
            BOOKS_EXCEL
        )

        sheet = workbook.active

        excel_row = None

        for r in range(
            2,
            sheet.max_row + 1
        ):

            excel_title = sheet.cell(
                row=r,
                column=1
            ).value

            excel_author = sheet.cell(
                row=r,
                column=2
            ).value

            excel_congress = sheet.cell(
                row=r,
                column=3
            ).value

            if (

                str(excel_title)
                ==
                str(book.book_title)

                and

                str(excel_author)
                ==
                str(book.author_translator)

                and

                str(excel_congress)
                ==
                str(book.congress_number)

            ):

                excel_row = r

                break

        if excel_row is not None:

            sheet.delete_rows(
                excel_row,
                1
            )

            workbook.save(
                BOOKS_EXCEL
            )

        book_title_for_log = book.book_title

        db.session.delete(book)

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin:

            add_action(

                "ادمین {} {} کتاب «{}» را حذف کرد".format(

                    admin.first_name,

                    admin.last_name,

                    book_title_for_log

                )

            )

        db.session.commit()

        flash(
            "کتاب با موفقیت حذف شد",
            "success"
        )

    except Exception as e:

        db.session.rollback()

        flash(
            "خطا در حذف کتاب: " + str(e),
            "danger"
        )

    return redirect(
        '/show_books'
    )


# =========================================================
# 100 کتاب برتر
# =========================================================

@app.route('/top_books')
def top_books():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    books = Books.query.order_by(
        Books.borrow_count.desc()
    ).limit(100).all()

    return render_template(
        'top_books.html',
        books=books
    )


# =========================================================
# EDIT BOOK
# =========================================================

@app.route(
    '/edit_book/<int:book_id>',
    methods=['GET', 'POST']
)
def edit_book(book_id):

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    book = Books.query.get(book_id)

    if not book:

        flash(
            "کتاب پیدا نشد",
            "danger"
        )

        return redirect('/show_books')

    if request.method == 'POST':

        old_title = book.book_title
        old_author = book.author_translator
        old_congress = book.congress_number

        book.book_title = request.form.get(
            'book_title'
        )

        book.author_translator = request.form.get(
            'author_translator'
        )

        book.congress_number = request.form.get(
            'congress_number'
        )

        book.shelf_number = request.form.get(
            'shelf_number'
        )

        book.row_number = request.form.get(
            'row_number'
        )

        book.comprehensive = (

            request.form.get(
                'comprehensive'
            ) == 'True'

        )

        try:

            workbook = openpyxl.load_workbook(
                BOOKS_EXCEL
            )

            sheet = workbook.active

            excel_row = None

            for r in range(
                2,
                sheet.max_row + 1
            ):

                title = sheet.cell(
                    row=r,
                    column=1
                ).value

                author = sheet.cell(
                    row=r,
                    column=2
                ).value

                congress = sheet.cell(
                    row=r,
                    column=3
                ).value

                if (

                    str(title) == str(old_title)

                    and

                    str(author) == str(old_author)

                    and

                    str(congress) == str(old_congress)

                ):

                    excel_row = r

                    break

            if excel_row is not None:

                values = [

                    book.book_title,

                    book.author_translator,

                    book.congress_number,

                    book.shelf_number,

                    book.row_number,

                    book.comprehensive

                ]

                for column, value in enumerate(

                    values,

                    start=1

                ):

                    sheet.cell(

                        row=excel_row,

                        column=column,

                        value=value

                    )

                workbook.save(
                    BOOKS_EXCEL
                )

            admin = Admins.query.filter_by(
                user_name=user
            ).first()

            if admin:

                add_action(

                    "ادمین {} {} کتاب «{}» را ویرایش کرد".format(

                        admin.first_name,

                        admin.last_name,

                        book.book_title

                    )

                )

            db.session.commit()

            flash(
                "اطلاعات کتاب با موفقیت ویرایش شد",
                "success"
            )

            return redirect(
                '/show_books'
            )

        except Exception as e:

            db.session.rollback()

            flash(
                "خطا در ویرایش کتاب: " + str(e),
                "danger"
            )

            return redirect(
                '/show_books'
            )

    return render_template(
        'edit_book.html',
        book=book
    )


# =========================================================
# GROUP IMPORT BOOKS
# =========================================================

@app.route(
    '/import_books',
    methods=['GET', 'POST']
)
def import_books():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    if request.method == 'GET':

        return render_template(
            'import_books.html'
        )

    file = request.files.get(
        'excel_file'
    )

    if not file or file.filename == "":

        flash(
            "لطفاً یک فایل Excel انتخاب کنید",
            "danger"
        )

        return redirect('/import_books')

    if not file.filename.lower().endswith('.xlsx'):

        flash(
            "فقط فایل xlsx قابل قبول است",
            "danger"
        )

        return redirect('/import_books')

    try:

        uploaded_workbook = openpyxl.load_workbook(
            file
        )

        uploaded_sheet = uploaded_workbook.active

    except Exception as e:

        flash(
            "فایل Excel قابل خواندن نیست: " + str(e),
            "danger"
        )

        return redirect('/import_books')

    try:

        main_workbook = openpyxl.load_workbook(
            BOOKS_EXCEL
        )

        main_sheet = main_workbook.active

    except Exception as e:

        flash(
            "Books.xlsx پیدا نشد: " + str(e),
            "danger"
        )

        return redirect('/import_books')

    success_count = 0
    error_count = 0
    errors = []

    last_real_row = 1

    for r in range(
        2,
        main_sheet.max_row + 1
    ):

        if any(

            main_sheet.cell(
                row=r,
                column=c
            ).value not in (None, "")

            for c in range(1, 7)

        ):

            last_real_row = r

    try:

        for excel_row_number, row in enumerate(

            uploaded_sheet.iter_rows(
                min_row=2,
                values_only=True
            ),

            start=2

        ):

            values = list(row)

            if not any(
                value not in (None, "")
                for value in values
            ):

                continue

            if len(values) < 6:

                error_count += 1

                errors.append(

                    "ردیف {}: تعداد ستون‌ها باید حداقل 6 باشد".format(
                        excel_row_number
                    )

                )

                continue

            values = values[:6]

            book_title = values[0]
            author_translator = values[1]
            congress_number = values[2]
            shelf_number = values[3]
            row_number = values[4]
            comprehensive = values[5]

            required = {

                "عنوان کتاب": book_title,
                "نویسنده / مترجم": author_translator,
                "شماره کنگره": congress_number,
                "شماره قفسه": shelf_number,
                "شماره ردیف": row_number

            }

            invalid = []

            for name, value in required.items():

                if value is None:

                    invalid.append(name)

                elif isinstance(value, str) and not value.strip():

                    invalid.append(name)

            if invalid:

                error_count += 1

                errors.append(

                    "ردیف {}: فیلدهای خالی: {}".format(

                        excel_row_number,

                        ", ".join(invalid)

                    )

                )

                continue

            book_title = str(book_title).strip()
            author_translator = str(author_translator).strip()
            congress_number = str(congress_number).strip()
            shelf_number = str(shelf_number).strip()
            row_number = str(row_number).strip()

            if isinstance(comprehensive, bool):

                comprehensive_value = comprehensive

            else:

                text = str(
                    comprehensive
                ).strip().lower()

                if text in [
                    "true",
                    "1",
                    "yes",
                    "بله",
                    "مرجع"
                ]:

                    comprehensive_value = True

                elif text in [
                    "false",
                    "0",
                    "no",
                    "خیر",
                    "غیرمرجع",
                    "",
                    "none"
                ]:

                    comprehensive_value = False

                else:

                    error_count += 1

                    errors.append(

                        "ردیف {}: مقدار مرجع نامعتبر است".format(
                            excel_row_number
                        )

                    )

                    continue

            book = Books(

                book_title=book_title,

                author_translator=author_translator,

                congress_number=congress_number,

                shelf_number=shelf_number,

                row_number=row_number,

                comprehensive=comprehensive_value,

                borrower_code="None",

                available=False

            )

            db.session.add(book)

            last_real_row += 1

            new_values = [

                book_title,

                author_translator,

                congress_number,

                shelf_number,

                row_number,

                comprehensive_value

            ]

            for column, value in enumerate(

                new_values,

                start=1

            ):

                main_sheet.cell(

                    row=last_real_row,

                    column=column,

                    value=value

                )

            success_count += 1

        # -------------------------------------------------
        # ثبت عملیات
        # -------------------------------------------------

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin and success_count > 0:

            add_action(

                "ادمین {} {} تعداد {} کتاب را به صورت گروهی اضافه کرد".format(

                    admin.first_name,

                    admin.last_name,

                    success_count

                )

            )

        db.session.commit()

        main_workbook.save(
            BOOKS_EXCEL
        )

    except Exception as e:

        db.session.rollback()

        flash(
            "خطا در ذخیره کتاب‌ها: " + str(e),
            "danger"
        )

        return redirect('/import_books')

    if success_count > 0:

        flash(

            "{} کتاب با موفقیت اضافه شد".format(
                success_count
            ),

            "success"

        )

    for error in errors:

        flash(
            error,
            "warning"
        )

    return redirect('/import_books')


# =========================================================
# ADD STUDENT
# =========================================================

@app.route(
    '/add_student',
    methods=['GET', 'POST']
)
def add_student():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    if request.method == 'POST':

        first_name = request.form.get(
            'first_name'
        )

        last_name = request.form.get(
            'last_name'
        )

        national_id_code = request.form.get(
            'national_id_code'
        )

        student_id_code = request.form.get(
            'student_id_code'
        )

        phone_number = request.form.get(
            'phone_number'
        )

        phone_number_two = request.form.get(
            'phone_number_two'
        )

        father_name = request.form.get(
            'father_name'
        )

        birthday = request.form.get(
            'birthday'
        )

        school_name = request.form.get(
            'school_name'
        )

        address = request.form.get(
            'address'
        )

        registration_date = date.today()

        expiration_date = registration_date.replace(

            year=registration_date.year + 1

        )

        try:

            student = Students(

                first_name=first_name,

                last_name=last_name,

                national_id_code=national_id_code,

                student_id_code=student_id_code,

                phone_number=phone_number,

                phone_number_two=phone_number_two,

                father_name=father_name,

                birthday=birthday,

                school_name=school_name,

                address=address,

                borrowing_capacity=3,

                registration_date=str(
                    registration_date
                ),

                expiration_date=str(
                    expiration_date
                ),

                opportunity=3,

                expired=False

            )

            db.session.add(student)

            db.session.flush()

            workbook = openpyxl.load_workbook(
                STUDENTS_EXCEL
            )

            sheet = workbook.active

            last_real_row = 1

            for r in range(
                2,
                sheet.max_row + 1
            ):

                if any(

                    sheet.cell(
                        row=r,
                        column=c
                    ).value not in (None, "")

                    for c in range(1, 15)

                ):

                    last_real_row = r

            new_row = last_real_row + 1

            values = [

                first_name,
                last_name,
                national_id_code,
                student_id_code,
                phone_number,
                phone_number_two,
                father_name,
                birthday,
                school_name,
                3,
                str(registration_date),
                str(expiration_date),
                3,
                address

            ]

            for column, value in enumerate(

                values,

                start=1

            ):

                sheet.cell(

                    row=new_row,

                    column=column,

                    value=value

                )

            workbook.save(
                STUDENTS_EXCEL
            )

            admin = Admins.query.filter_by(
                user_name=user
            ).first()

            if admin:

                add_action(

                    "ادمین {} {} دانش‌آموز {} {} را اضافه کرد".format(

                        admin.first_name,

                        admin.last_name,

                        student.first_name,

                        student.last_name

                    )

                )

            db.session.commit()

            flash(
                "دانش‌آموز با موفقیت ثبت شد",
                "success"
            )

            return redirect(
                '/show_students'
            )

        except Exception as e:

            db.session.rollback()

            flash(
                "خطا در ثبت دانش‌آموز: " + str(e),
                "danger"
            )

            return redirect('/add_student')

    return render_template(
        'add_student.html'
    )


# =========================================================
# SHOW STUDENTS
# =========================================================

@app.route('/show_students')
def show_students():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    search = request.args.get(
        "search",
        ""
    ).strip()

    page = request.args.get(
        "page",
        1,
        type=int
    )

    if search:

        students = Students.query.filter(

            (Students.first_name.contains(search)) |

            (Students.last_name.contains(search)) |

            (Students.national_id_code.contains(search)) |

            (Students.student_id_code.contains(search)) |

            (Students.phone_number.contains(search)) |

            (Students.school_name.contains(search))

        ).paginate(

            page=page,

            per_page=30,

            error_out=False

        )

    else:

        students = Students.query.paginate(

            page=page,

            per_page=30,

            error_out=False

        )

    return render_template(

        'students.html',

        students=students,

        search=search

    )


# =========================================================
# GROUP IMPORT STUDENTS
#
# فایل ورودی:
#
# A نام
# B نام خانوادگی
# C کد ملی
# D کد دانش‌آموزی
# E شماره تلفن
# F شماره تلفن دوم
# G نام پدر
# H تاریخ تولد
# I نام مدرسه
# J آدرس
#
# فایل اصلی:
#
# A نام
# B نام خانوادگی
# C کد ملی
# D کد دانش‌آموزی
# E شماره تلفن
# F شماره تلفن دوم
# G نام پدر
# H تاریخ تولد
# I نام مدرسه
# J ظرفیت
# K تاریخ ثبت
# L تاریخ انقضا
# M فرصت
# N آدرس
# =========================================================

@app.route(
    '/import_students',
    methods=['GET', 'POST']
)
def import_students():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    if request.method == 'GET':

        return render_template(
            'import_students.html'
        )

    file = request.files.get(
        'excel_file'
    )

    if not file or file.filename == "":

        flash(
            "لطفاً یک فایل Excel انتخاب کنید",
            "danger"
        )

        return redirect('/import_students')

    if not file.filename.lower().endswith('.xlsx'):

        flash(
            "فایل باید xlsx باشد",
            "danger"
        )

        return redirect('/import_students')

    try:

        uploaded_workbook = openpyxl.load_workbook(
            file
        )

        uploaded_sheet = uploaded_workbook.active

    except Exception as e:

        flash(
            "فایل Excel قابل خواندن نیست: " + str(e),
            "danger"
        )

        return redirect('/import_students')

    try:

        main_workbook = openpyxl.load_workbook(
            STUDENTS_EXCEL
        )

        main_sheet = main_workbook.active

    except Exception as e:

        flash(
            "Students.xlsx پیدا نشد: " + str(e),
            "danger"
        )

        return redirect('/import_students')

    registration_date = date.today()

    expiration_date = registration_date.replace(

        year=registration_date.year + 1

    )

    students_to_add = []

    errors = []

    # -----------------------------------------------------
    # خواندن Excel
    # -----------------------------------------------------

    for row_number in range(
        2,
        uploaded_sheet.max_row + 1
    ):

        values = [

            uploaded_sheet.cell(

                row=row_number,

                column=column

            ).value

            for column in range(1, 11)

        ]

        if all(
            value in (None, "")
            for value in values
        ):

            continue

        first_name = values[0]
        last_name = values[1]
        national_id_code = values[2]
        student_id_code = values[3]
        phone_number = values[4]
        phone_number_two = values[5]
        father_name = values[6]
        birthday = values[7]
        school_name = values[8]
        address = values[9]

        required = {

            "نام": first_name,

            "نام خانوادگی": last_name,

            "کد ملی": national_id_code,

            "کد دانش‌آموزی": student_id_code,

            "شماره تلفن": phone_number

        }

        invalid = []

        for name, value in required.items():

            if value is None:

                invalid.append(name)

            elif isinstance(value, str):

                if not value.strip():

                    invalid.append(name)

        if invalid:

            errors.append(

                "ردیف {}: فیلدهای خالی: {}".format(

                    row_number,

                    ", ".join(invalid)

                )

            )

            continue

        first_name = str(first_name).strip()
        last_name = str(last_name).strip()
        national_id_code = str(national_id_code).strip()
        student_id_code = str(student_id_code).strip()
        phone_number = str(phone_number).strip()

        if phone_number_two not in (None, ""):

            phone_number_two = str(
                phone_number_two
            ).strip()

        else:

            phone_number_two = None

        if father_name not in (None, ""):

            father_name = str(
                father_name
            ).strip()

        else:

            father_name = None

        if birthday not in (None, ""):

            birthday = str(
                birthday
            ).strip()

        else:

            birthday = None

        if school_name not in (None, ""):

            school_name = str(
                school_name
            ).strip()

        else:

            school_name = None

        # -------------------------------------------------
        # آدرس ستون J
        # -------------------------------------------------

        if address not in (None, ""):

            address = str(
                address
            ).strip()

        else:

            address = None

        borrowing_capacity = 3

        excel_registration_date = str(
            registration_date
        )

        excel_expiration_date = str(
            expiration_date
        )

        opportunity = 3

        student = Students(

            first_name=first_name,

            last_name=last_name,

            national_id_code=national_id_code,

            student_id_code=student_id_code,

            phone_number=phone_number,

            phone_number_two=phone_number_two,

            father_name=father_name,

            birthday=birthday,

            school_name=school_name,

            borrowing_capacity=borrowing_capacity,

            registration_date=excel_registration_date,

            expiration_date=excel_expiration_date,

            opportunity=opportunity,

            expired=False,

            address=address

        )

        excel_values = [

            first_name,

            last_name,

            national_id_code,

            student_id_code,

            phone_number,

            phone_number_two,

            father_name,

            birthday,

            school_name,

            borrowing_capacity,

            excel_registration_date,

            excel_expiration_date,

            opportunity,

            address

        ]

        students_to_add.append(

            (
                student,
                excel_values
            )

        )

    if errors:

        for error in errors:

            flash(
                error,
                "warning"
            )

        return redirect('/import_students')

    if not students_to_add:

        flash(
            "هیچ دانش‌آموز معتبری پیدا نشد",
            "warning"
        )

        return redirect('/import_students')

    # -----------------------------------------------------
    # آخرین ردیف واقعی
    # -----------------------------------------------------

    last_real_row = 1

    for r in range(
        2,
        main_sheet.max_row + 1
    ):

        if any(

            main_sheet.cell(
                row=r,
                column=c
            ).value not in (None, "")

            for c in range(1, 15)

        ):

            last_real_row = r

    try:

        for student, excel_values in students_to_add:

            db.session.add(student)

            last_real_row += 1

            for column, value in enumerate(

                excel_values,

                start=1

            ):

                main_sheet.cell(

                    row=last_real_row,

                    column=column,

                    value=value

                )

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin:

            add_action(

                "ادمین {} {} تعداد {} دانش‌آموز را به صورت گروهی اضافه کرد".format(

                    admin.first_name,

                    admin.last_name,

                    len(students_to_add)

                )

            )

        db.session.commit()

        main_workbook.save(
            STUDENTS_EXCEL
        )

        flash(

            "{} دانش‌آموز با موفقیت اضافه شدند".format(
                len(students_to_add)
            ),

            "success"

        )

        return redirect('/show_students')

    except Exception as e:

        db.session.rollback()

        flash(

            "خطا در افزودن گروهی دانش‌آموزان: " + str(e),

            "danger"

        )

        return redirect('/import_students')


# =========================================================
# EDIT STUDENT
# =========================================================

@app.route(
    '/edit_student/<int:student_id>',
    methods=['GET', 'POST']
)
def edit_student(student_id):

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    student = Students.query.get_or_404(
        student_id
    )

    if request.method == 'POST':

        old_national_id = student.national_id_code
        old_student_id = student.student_id_code

        student.first_name = request.form.get(
            'first_name'
        )

        student.last_name = request.form.get(
            'last_name'
        )

        student.national_id_code = request.form.get(
            'national_id_code'
        )

        student.student_id_code = request.form.get(
            'student_id_code'
        )

        student.phone_number = request.form.get(
            'phone_number'
        )

        student.phone_number_two = request.form.get(
            'phone_number_two'
        )

        student.father_name = request.form.get(
            'father_name'
        )

        student.birthday = request.form.get(
            'birthday'
        )

        student.school_name = request.form.get(
            'school_name'
        )

        student.address = request.form.get(
            'address'
        )

        try:

            workbook = openpyxl.load_workbook(
                STUDENTS_EXCEL
            )

            sheet = workbook.active

            found_in_excel = False

            for row in sheet.iter_rows(
                min_row=2
            ):

                excel_national_id = row[2].value
                excel_student_id = row[3].value

                if (

                    str(excel_national_id).strip()
                    ==
                    str(old_national_id).strip()

                    and

                    str(excel_student_id).strip()
                    ==
                    str(old_student_id).strip()

                ):

                    row[0].value = student.first_name
                    row[1].value = student.last_name
                    row[2].value = student.national_id_code
                    row[3].value = student.student_id_code
                    row[4].value = student.phone_number
                    row[5].value = student.phone_number_two
                    row[6].value = student.father_name
                    row[7].value = student.birthday
                    row[8].value = student.school_name
                    row[9].value = student.borrowing_capacity
                    row[10].value = student.registration_date
                    row[11].value = student.expiration_date
                    row[12].value = student.opportunity
                    row[13].value = student.address

                    found_in_excel = True

                    break

            if not found_in_excel:

                db.session.rollback()

                flash(

                    "دانش‌آموز در فایل Students.xlsx پیدا نشد",

                    "danger"

                )

                return redirect(
                    '/show_students'
                )

            workbook.save(
                STUDENTS_EXCEL
            )

            admin = Admins.query.filter_by(
                user_name=user
            ).first()

            if admin:

                add_action(

                    "ادمین {} {} دانش‌آموز {} {} را ویرایش کرد".format(

                        admin.first_name,

                        admin.last_name,

                        student.first_name,

                        student.last_name

                    )

                )

            db.session.commit()

            flash(
                "اطلاعات دانش‌آموز با موفقیت ویرایش شد",
                "success"
            )

            return redirect('/show_students')

        except Exception as e:

            db.session.rollback()

            flash(

                "خطا در ویرایش دانش‌آموز: " + str(e),

                "danger"

            )

            return redirect(
                '/show_students'
            )

    return render_template(
        'edit_student.html',
        student=student
    )


# =========================================================
# DELETE STUDENT
# =========================================================

@app.route('/delete_student/<int:student_id>')
def delete_student(student_id):

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    student = Students.query.get_or_404(
        student_id
    )

    try:

        workbook = openpyxl.load_workbook(
            STUDENTS_EXCEL
        )

        sheet = workbook.active

        found_in_excel = False

        for row in sheet.iter_rows(
            min_row=2
        ):

            excel_national_id = row[2].value
            excel_student_id = row[3].value

            if (

                str(excel_national_id).strip()
                ==
                str(student.national_id_code).strip()

                and

                str(excel_student_id).strip()
                ==
                str(student.student_id_code).strip()

            ):

                sheet.delete_rows(
                    row[0].row,
                    1
                )

                found_in_excel = True

                break

        if not found_in_excel:

            flash(

                "دانش‌آموز در فایل Students.xlsx پیدا نشد",

                "danger"

            )

            return redirect('/show_students')

        workbook.save(
            STUDENTS_EXCEL
        )

        first_name = student.first_name
        last_name = student.last_name

        db.session.delete(student)

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin:

            add_action(

                "ادمین {} {} دانش‌آموز {} {} را حذف کرد".format(

                    admin.first_name,

                    admin.last_name,

                    first_name,

                    last_name

                )

            )

        db.session.commit()

        flash(
            "دانش‌آموز با موفقیت حذف شد",
            "success"
        )

    except Exception as e:

        db.session.rollback()

        flash(
            "خطا در حذف دانش‌آموز: " + str(e),
            "danger"
        )

    return redirect('/show_students')


# =========================================================
# صفحه ثبت امانت
# =========================================================

@app.route('/new_borrowing')
def new_borrowing():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    return render_template(
        'new_borrowing.html'
    )


# =========================================================
# جستجوی دانش‌آموز
# =========================================================

@app.route('/search_students')
def search_students():

    user = request.cookies.get("user")

    if not user:

        return jsonify([])

    search = request.args.get(
        'search',
        ''
    ).strip()

    if not search:

        return jsonify([])

    students = Students.query.filter(

        (Students.first_name.contains(search)) |

        (Students.last_name.contains(search)) |

        (Students.student_id_code.contains(search))

    ).limit(10).all()

    result = []

    for student in students:

        result.append({

            'id':
                student.id,

            'first_name':
                student.first_name or '',

            'last_name':
                student.last_name or '',

            'student_id_code':
                student.student_id_code or ''

        })

    return jsonify(result)


# =========================================================
# جستجوی کتاب
# =========================================================

@app.route('/search_books')
def search_books():

    user = request.cookies.get("user")

    if not user:

        return jsonify([])

    search = request.args.get(
        'search',
        ''
    ).strip()

    if not search:

        return jsonify([])

    books = Books.query.filter(

        (Books.book_title.contains(search)) |

        (Books.congress_number.contains(search))

    ).limit(15).all()

    result = []

    for book in books:

        if book.available:

            continue

        result.append({

            'id':
                book.id,

            'book_title':
                book.book_title or '',

            'congress_number':
                book.congress_number or '',

            'comprehensive':
                bool(book.comprehensive)

        })

    return jsonify(result)


# =========================================================
# ثبت نهایی امانت
# =========================================================

@app.route(
    '/save_borrowing',
    methods=['POST']
)
def save_borrowing():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    try:

        # =================================================
        # اطلاعات فرم
        # =================================================

        student_id = request.form.get(
            'student_id'
        )

        book1_id = request.form.get(
            'book1_id'
        )

        book2_id = request.form.get(
            'book2_id'
        )

        book3_id = request.form.get(
            'book3_id'
        )

        # =================================================
        # بررسی دانش‌آموز
        # =================================================

        if not student_id:

            flash(
                "لطفاً یک دانش‌آموز انتخاب کنید.",
                "danger"
            )

            return redirect('/new_borrowing')

        # -------------------------------------------------
        # اینجا مشکل اصلی قبلی اصلاح شده است.
        # query.get فقط یک id می‌گیرد.
        # -------------------------------------------------

        student = Students.query.get(
            student_id
        )

        if not student:

            flash(
                "دانش‌آموز پیدا نشد.",
                "danger"
            )

            return redirect('/new_borrowing')

        # =================================================
        # بررسی وضعیت دانش‌آموز
        # =================================================

        # دوباره همین لحظه بررسی می‌کنیم تا اگر
        # فرصت دانش‌آموز در همین Request تمام شده باشد
        # اجازه امانت داده نشود.

        if student.opportunity is not None:

            if student.opportunity <= 0:

                student.opportunity = 0

                student.expired = True

                db.session.commit()

        if student.expired:

            flash(

                "کارت این دانش‌آموز منقضی شده است و امکان امانت گرفتن کتاب ندارد.",

                "danger"

            )

            return redirect('/new_borrowing')

        # =================================================
        # کتاب‌ها
        # =================================================

        book_ids = [

            book1_id,

            book2_id,

            book3_id

        ]

        book_ids = [

            book_id

            for book_id in book_ids

            if book_id

        ]

        # =================================================
        # حداقل یک کتاب
        # =================================================

        if not book_ids:

            flash(
                "لطفاً حداقل یک کتاب انتخاب کنید.",
                "danger"
            )

            return redirect('/new_borrowing')

        # =================================================
        # جلوگیری از تکراری بودن
        # =================================================

        if len(book_ids) != len(set(book_ids)):

            flash(

                "یک کتاب را نمی‌توان بیشتر از یک بار انتخاب کرد.",

                "danger"

            )

            return redirect('/new_borrowing')

        # =================================================
        # ظرفیت
        # =================================================

        if student.borrowing_capacity is None:

            student.borrowing_capacity = 0

        if student.borrowing_capacity <= 0:

            flash(

                "ظرفیت امانت این دانش‌آموز تمام شده است.",

                "danger"

            )

            return redirect('/new_borrowing')

        if len(book_ids) > student.borrowing_capacity:

            flash(

                "ظرفیت باقی‌مانده امانت این دانش‌آموز {} کتاب است.".format(

                    student.borrowing_capacity

                ),

                "danger"

            )

            return redirect('/new_borrowing')

        # =================================================
        # پیدا کردن کتاب‌ها
        # =================================================

        selected_books = []

        for book_id in book_ids:

            book = Books.query.get(
                book_id
            )

            if not book:

                flash(

                    "یکی از کتاب‌های انتخاب شده پیدا نشد.",

                    "danger"

                )

                return redirect('/new_borrowing')

            if book.comprehensive:

                flash(

                    "کتاب «{}» مرجع است و امکان امانت دادن آن وجود ندارد.".format(

                        book.book_title

                    ),

                    "danger"

                )

                return redirect('/new_borrowing')

            if book.available:

                flash(

                    "کتاب «{}» قبلاً امانت داده شده است.".format(

                        book.book_title

                    ),

                    "danger"

                )

                return redirect('/new_borrowing')

            selected_books.append(book)

        # =================================================
        # تاریخ امانت و بازگشت
        # =================================================

        now = datetime.now()

        borrowing_date = now.strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        return_date = (

            now + timedelta(days=14)

        ).strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        # =================================================
        # ثبت امانت
        # =================================================

        for book in selected_books:

            borrowing = Borrowing(

                student_id_code_borrow=
                    student.student_id_code,

                book_id_borrow=
                    str(book.id),

                date_of_borrowing=
                    borrowing_date,

                return_date=
                    return_date,

                extension=False,

                late_penalty_applied=False

            )

            db.session.add(
                borrowing
            )

            book.borrower_code = (
                student.student_id_code
            )

            book.available = True

            book.borrow_count = (
                book.borrow_count or 0
            ) + 1

        # =================================================
        # کم کردن ظرفیت
        # =================================================

        student.borrowing_capacity = (

            student.borrowing_capacity
            -
            len(selected_books)

        )

        # =================================================
        # ثبت عملیات
        # =================================================

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin:

            add_action(

                "ادمین {} {} تعداد {} کتاب را به {} {} امانت داد".format(

                    admin.first_name,

                    admin.last_name,

                    len(selected_books),

                    student.first_name,

                    student.last_name

                )

            )

        # =================================================
        # ذخیره
        # =================================================

        db.session.commit()

        # =================================================
        # پیام‌ها
        # =================================================

        for book in selected_books:

            flash(

                "کتاب «{}» با موفقیت به {} {} امانت داده شد.".format(

                    book.book_title,

                    student.first_name,

                    student.last_name

                ),

                "success"

            )

        flash(

            "ظرفیت باقی‌مانده امانت: {} کتاب".format(

                student.borrowing_capacity

            ),

            "success"

        )

        flash(

            "تاریخ بازگشت کتاب‌ها: {}".format(
                return_date
            ),

            "success"

        )

        return redirect('/new_borrowing')

    except Exception as e:

        db.session.rollback()

        flash(

            "خطا در ثبت امانت: {}".format(e),

            "danger"

        )

        return redirect('/new_borrowing')


# =========================================================
# SHOW BORROWINGS
# =========================================================

@app.route('/show_borrowings')
def show_borrowings():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    search = request.args.get(
        'search',
        ''
    ).strip()

    borrowings = Borrowing.query.order_by(
        Borrowing.id.desc()
    ).all()

    borrowing_data = []

    now = datetime.now()

    for borrowing in borrowings:

        try:

            book = Books.query.get(
                int(borrowing.book_id_borrow)
            )

        except (ValueError, TypeError):

            continue

        student = Students.query.filter_by(

            student_id_code=
                borrowing.student_id_code_borrow

        ).first()

        if not book or not student:

            continue

        # -------------------------------------------------
        # دیرکرد
        # -------------------------------------------------

        return_date = parse_datetime(
            borrowing.return_date
        )

        expired = (

            return_date is not None
            and
            return_date < now

        )

        # -------------------------------------------------
        # جستجو
        # -------------------------------------------------

        if search:

            search_lower = search.lower()

            book_title = str(
                book.book_title or ''
            ).lower()

            congress_number = str(
                book.congress_number or ''
            ).lower()

            first_name = str(
                student.first_name or ''
            ).lower()

            last_name = str(
                student.last_name or ''
            ).lower()

            student_id_code = str(
                student.student_id_code or ''
            ).lower()

            full_name = (
                first_name
                + ' '
                + last_name
            )

            if not (

                search_lower in book_title

                or

                search_lower in congress_number

                or

                search_lower in first_name

                or

                search_lower in last_name

                or

                search_lower in full_name

                or

                search_lower in student_id_code

            ):

                continue

        borrowing_data.append({

            "borrowing":
                borrowing,

            "book":
                book,

            "student":
                student,

            "expired":
                expired

        })

    # =====================================================
    # Pagination
    # =====================================================

    page = request.args.get(
        'page',
        1,
        type=int
    )

    if page < 1:

        page = 1

    per_page = 30

    total = len(borrowing_data)

    total_pages = (

        (total + per_page - 1)
        // per_page

        if total > 0

        else 1

    )

    if page > total_pages:

        page = total_pages

    start = (
        page - 1
    ) * per_page

    end = start + per_page

    page_items = borrowing_data[
        start:end
    ]

    return render_template(

        'show_borrowings.html',

        borrowing_data=page_items,

        page=page,

        total_pages=total_pages,

        search=search

    )


# =========================================================
# EXTEND BORROWING
# =========================================================

@app.route('/extend_borrowing/<int:borrowing_id>')
def extend_borrowing(borrowing_id):

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    borrowing = Borrowing.query.get_or_404(
        borrowing_id
    )

    if borrowing.extension:

        flash(

            "این امانت قبلاً تمدید شده است.",

            "danger"

        )

        return redirect('/show_borrowings')

    try:

        old_return_date = parse_datetime(
            borrowing.return_date
        )

        if not old_return_date:

            flash(
                "تاریخ بازگشت امانت نامعتبر است.",
                "danger"
            )

            return redirect('/show_borrowings')

        new_return_date = (

            old_return_date
            +
            timedelta(days=14)

        )

        borrowing.return_date = (
            new_return_date.strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )

        borrowing.extension = True

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin:

            try:

                book = Books.query.get(

                    int(
                        borrowing.book_id_borrow
                    )

                )

            except (ValueError, TypeError):

                book = None

            student = Students.query.filter_by(

                student_id_code=
                    borrowing.student_id_code_borrow

            ).first()

            if book and student:

                add_action(

                    "ادمین {} {} امانت کتاب «{}» برای دانش‌آموز {} {} را ۱۴ روز تمدید کرد".format(

                        admin.first_name,

                        admin.last_name,

                        book.book_title,

                        student.first_name,

                        student.last_name

                    )

                )

        db.session.commit()

        flash(

            "امانت با موفقیت ۱۴ روز تمدید شد.",

            "success"

        )

        return redirect('/show_borrowings')

    except Exception as e:

        db.session.rollback()

        flash(

            "خطا در تمدید امانت: {}".format(e),

            "danger"

        )

        return redirect('/show_borrowings')


# =========================================================
# RETURN BORROWING
#
# مهم:
# اینجا opportunity کم نمی‌شود.
#
# دیرکرد قبلاً در before_request بررسی شده است.
# =========================================================

@app.route('/return_borrowing/<int:borrowing_id>')
def return_borrowing(borrowing_id):

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    borrowing = Borrowing.query.get_or_404(
        borrowing_id
    )

    try:

        # -------------------------------------------------
        # کتاب
        # -------------------------------------------------

        try:

            book = Books.query.get(

                int(
                    borrowing.book_id_borrow
                )

            )

        except (ValueError, TypeError):

            book = None

        # -------------------------------------------------
        # دانش‌آموز
        # -------------------------------------------------

        student = Students.query.filter_by(

            student_id_code=
                borrowing.student_id_code_borrow

        ).first()

        # -------------------------------------------------
        # آزاد کردن کتاب
        # -------------------------------------------------

        if book:

            book.available = False

            book.borrower_code = "None"

        # -------------------------------------------------
        # برگرداندن ظرفیت
        # -------------------------------------------------

        if student:

            if student.borrowing_capacity is None:

                student.borrowing_capacity = 0

            student.borrowing_capacity += 1

        # -------------------------------------------------
        # اطلاعات عملیات
        # -------------------------------------------------

        book_title_for_log = (

            book.book_title

            if book

            else "نامشخص"

        )

        student_name_for_log = (
            "دانش‌آموز نامشخص"
        )

        if student:

            student_name_for_log = "{} {}".format(

                student.first_name or "",

                student.last_name or ""

            ).strip()

        # -------------------------------------------------
        # حذف امانت
        # -------------------------------------------------

        db.session.delete(
            borrowing
        )

        # -------------------------------------------------
        # ثبت عملیات
        # -------------------------------------------------

        admin = Admins.query.filter_by(
            user_name=user
        ).first()

        if admin:

            add_action(

                "ادمین {} {} کتاب «{}» را از {} پس گرفت".format(

                    admin.first_name,

                    admin.last_name,

                    book_title_for_log,

                    student_name_for_log

                )

            )

        # -------------------------------------------------
        # ذخیره
        # -------------------------------------------------

        db.session.commit()

        flash(

            "کتاب با موفقیت پس گرفته شد.",

            "success"

        )

        return redirect('/show_borrowings')

    except Exception as e:

        db.session.rollback()

        flash(

            "خطا در پس گرفتن کتاب: {}".format(e),

            "danger"

        )

        return redirect('/show_borrowings')


# =========================================================
# OVERDUE BORROWINGS
# =========================================================

@app.route('/overdue_borrowings')
def overdue_borrowings():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    search = request.args.get(
        'search',
        ''
    ).strip()

    borrowings = Borrowing.query.order_by(
        Borrowing.id.desc()
    ).all()

    borrowing_data = []

    now = datetime.now()

    for borrowing in borrowings:

        return_date = parse_datetime(
            borrowing.return_date
        )

        if not return_date:

            continue

        if return_date >= now:

            continue

        try:

            book = Books.query.get(

                int(
                    borrowing.book_id_borrow
                )

            )

        except (ValueError, TypeError):

            continue

        student = Students.query.filter_by(

            student_id_code=
                borrowing.student_id_code_borrow

        ).first()

        if not book or not student:

            continue

        # -------------------------------------------------
        # جستجو
        # -------------------------------------------------

        if search:

            search_lower = search.lower()

            book_title = str(
                book.book_title or ''
            ).lower()

            congress_number = str(
                book.congress_number or ''
            ).lower()

            first_name = str(
                student.first_name or ''
            ).lower()

            last_name = str(
                student.last_name or ''
            ).lower()

            student_id_code = str(
                student.student_id_code or ''
            ).lower()

            full_name = (
                first_name
                + ' '
                + last_name
            )

            if not (

                search_lower in book_title

                or

                search_lower in congress_number

                or

                search_lower in first_name

                or

                search_lower in last_name

                or

                search_lower in full_name

                or

                search_lower in student_id_code

            ):

                continue

        borrowing_data.append({

            "borrowing":
                borrowing,

            "book":
                book,

            "student":
                student

        })

    # =====================================================
    # Pagination
    # =====================================================

    page = request.args.get(
        'page',
        1,
        type=int
    )

    if page < 1:

        page = 1

    per_page = 30

    total = len(borrowing_data)

    total_pages = (

        (total + per_page - 1)
        // per_page

        if total > 0

        else 1

    )

    if page > total_pages:

        page = total_pages

    start = (
        page - 1
    ) * per_page

    end = start + per_page

    page_items = borrowing_data[
        start:end
    ]

    return render_template(

        'overdue_borrowings.html',

        borrowing_data=page_items,

        page=page,

        total_pages=total_pages,

        search=search

    )


# =========================================================
# SHOW ACTIONS
# =========================================================

@app.route('/show_actions')
def show_actions():

    user = request.cookies.get("user")

    if not user:

        return redirect('/login')

    current_admin = Admins.query.filter_by(
        user_name=user
    ).first()

    if not current_admin or not current_admin.god_user:

        flash(

            "شما اجازه دسترسی به صفحه گزارش عملیات را ندارید",

            "danger"

        )

        return redirect('/')

    search = request.args.get(
        'search',
        ''
    ).strip()

    actions = Actions.query.order_by(
        Actions.id.desc()
    ).all()

    if search:

        search_lower = search.lower()

        filtered_actions = []

        for action in actions:

            action_text = str(
                action.actions_taken or ''
            ).lower()

            action_date = str(
                action.date or ''
            ).lower()

            action_time = str(
                action.time or ''
            ).lower()

            if (

                search_lower in action_text

                or

                search_lower in action_date

                or

                search_lower in action_time

            ):

                filtered_actions.append(
                    action
                )

        actions = filtered_actions

    # =====================================================
    # Pagination
    # =====================================================

    page = request.args.get(
        'page',
        1,
        type=int
    )

    if page < 1:

        page = 1

    per_page = 30

    total = len(actions)

    total_pages = (

        (total + per_page - 1)
        // per_page

        if total > 0

        else 1

    )

    if page > total_pages:

        page = total_pages

    start = (
        page - 1
    ) * per_page

    end = start + per_page

    page_actions = actions[
        start:end
    ]

    return render_template(

        'show_actions.html',

        actions=page_actions,

        page=page,

        total_pages=total_pages,

        search=search

    )






@app.route("/expired_students")
def expired_students():

    username = request.cookies.get("user")

    if not username:
        return redirect("/login")

    admin = Admins.query.filter_by(
        user_name=username
    ).first()

    if not admin:
        return redirect("/login")

    # -----------------------------
    # جستجو
    # -----------------------------

    search = request.args.get(
        "search",
        ""
    ).strip()

    # -----------------------------
    # شماره صفحه
    # -----------------------------

    page = request.args.get(
        "page",
        1,
        type=int
    )

    # -----------------------------
    # فقط دانش‌آموزان منقضی
    # -----------------------------

    query = Students.query.filter_by(
        expired=True
    )

    # -----------------------------
    # جستجو
    # -----------------------------

    if search:

        query = query.filter(
            db.or_(
                Students.first_name.contains(search),
                Students.last_name.contains(search),
                Students.national_id_code.contains(search),
                Students.student_id_code.contains(search),
                Students.school_name.contains(search)
            )
        )

    # -----------------------------
    # صفحه‌بندی 30 تایی
    # -----------------------------

    students = query.order_by(
        Students.first_name.asc()
    ).paginate(
        page=page,
        per_page=30,
        error_out=False
    )

    return render_template(
        "expired_students.html",
        students=students,
        search=search
    )


@app.route("/renew_student/<int:student_id>")
def renew_student(student_id):

    # بررسی ورود ادمین
    username = request.cookies.get("user")

    if not username:
        return redirect("/login")

    admin = Admins.query.filter_by(user_name=username).first()

    if not admin:
        return redirect("/login")

    # پیدا کردن دانش‌آموز
    student = db.session.get(Students, student_id)

    if not student:
        flash("دانش‌آموز پیدا نشد.", "error")
        return redirect("/expired_students")

    # فقط دانش‌آموز منقضی شده قابلیت تمدید دارد
    if not student.expired:
        flash("کارت این دانش‌آموز هنوز منقضی نشده است.", "warning")
        return redirect("/expired_students")

    try:

        # =====================================================
        # 1. پیدا کردن تمام امانت‌های فعال دانش‌آموز
        # =====================================================

        borrowings = Borrowing.query.filter_by(
            student_id_code_borrow=student.student_id_code
        ).all()


        # =====================================================
        # 2. پس گرفتن تمام کتاب‌های دانش‌آموز
        # =====================================================

        for borrowing in borrowings:

            book = db.session.get(
                Books,
                borrowing.book_id_borrow
            )

            if book:

                # آزاد کردن کتاب
                book.available = False

                # حذف نام/کد امانت‌گیرنده
                book.borrower_code = "None"


            # حذف رکورد امانت
            db.session.delete(borrowing)


        # =====================================================
        # 3. تمدید کارت
        # =====================================================

        today = date.today()

        next_year = today.replace(
            year=today.year + 1
        )

        student.registration_date = str(today)

        student.expiration_date = str(next_year)

        # فرصت‌های دانش‌آموز از اول
        student.opportunity = 3

        # ظرفیت امانت دوباره 3
        student.borrowing_capacity = 3

        # کارت فعال
        student.expired = False


        # =====================================================
        # 4. ثبت تغییرات
        # =====================================================

        add_action(
            "کارت دانش‌آموز "
            + str(student.first_name)
            + " "
            + str(student.last_name)
            + " تمدید شد و تمام امانت‌های قبلی او بسته شدند."
        )


        db.session.commit()

        flash(
            "کارت دانش‌آموز با موفقیت تمدید شد و امانت‌های قبلی او بسته شدند.",
            "success"
        )

    except Exception as e:

        db.session.rollback()

        print("خطا در تمدید کارت:", e)

        flash(
            "خطا در تمدید کارت: " + str(e),
            "error"
        )


    return redirect("/expired_students")





# =========================================================
# ERROR 500
# =========================================================

@app.errorhandler(500)
def internal_server_error(error):

    return render_template(
        "error500.html"
    ), 500


# =========================================================
# RUN
# =========================================================

if __name__ == '__main__':

    app.run(

        host='0.0.0.0',

        port=4540

    )